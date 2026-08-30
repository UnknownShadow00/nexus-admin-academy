"""Safe, certification-agnostic intake for externally approved V2 curriculum.

The inbox is not a second curriculum engine.  This service validates and
normalizes an author-friendly package into the files consumed by the existing
V2 loaders.  Educational prose and answers are never rewritten.
"""

from __future__ import annotations

import csv
import hashlib
import json
import os
import random
import re
import shutil
import stat
import subprocess
import sys
import tempfile
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath, PureWindowsPath
from typing import Callable
from zipfile import BadZipFile, ZipFile

import openpyxl
import yaml

from app.services.question_importer import (
    TEMPLATE_COLUMNS,
    objective_codes as parse_objective_codes,
    parse_csv_file,
    parse_xlsx_file,
    row_to_payload,
)
from app.services.question_validation import validate_question
from app.services.v2_assessment_selector import ConstraintSelectionError, select_constrained

MAX_ZIP_FILES = 2_000
MAX_ZIP_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
APPROVALS_FILE = "editorial-approvals.yaml"

FIELD_ALIASES = {
    "type": "question_type",
    "question": "question_text",
    "prompt": "question_text",
    "correct_answer": "correct_answers",
    "objective": "objective_code",
    "objectives": "objective_code",
    "lesson_key": "lesson_id",
    "accepted_variants": "acceptable_answers",
    "minimum_concepts": "min_concepts_for_pass",
    "min_concepts_pass": "min_concepts_for_pass",
    "provenance": "source_name",
    "provenance_source": "source_name",
    "editorial_status": "final_validation_status",
}
QUESTION_TYPE_ALIASES = {
    "single-choice": "single",
    "single_choice": "single",
    "multiple-choice": "single",
    "multi-select": "multi",
    "multi_select": "multi",
    "short-answer": "short_answer",
    "free-response": "free_response",
    "true-false": "true_false",
}
EDITORIAL_STATUS_ALIASES = {
    "approve": "APPROVED",
    "approved": "APPROVED",
    "edit": "APPROVED_AFTER_EDIT",
    "approved_after_edit": "APPROVED_AFTER_EDIT",
}
IMPORTANCE_ALIASES = {
    "job critical": "job_critical",
    "job-critical": "job_critical",
    "working knowledge": "working_knowledge",
    "working-knowledge": "working_knowledge",
    "awareness": "awareness",
    "know_it": "working_knowledge",
}
RELATIONSHIP_ALIASES = {
    "new": "new",
    "review": "review",
    "review of": "review",
    "review_of": "review",
    "deep dive": "deep_dive",
    "deep-dive": "deep_dive",
    "deep_dive": "deep_dive",
}


class IntakeError(ValueError):
    """A package-local validation error suitable for an owner-facing report."""

    def __init__(self, message: str, *, file: str = "package", field: str = "package"):
        super().__init__(message)
        self.file = file
        self.field = field

    def as_dict(self) -> dict:
        return {
            "file": self.file,
            "row": None,
            "key": None,
            "field": self.field,
            "expected": None,
            "actual": None,
            "message": str(self),
        }


def _within(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
        return True
    except ValueError:
        return False


def safe_extract_zip(archive_path: Path, destination: Path) -> None:
    """Extract a ZIP only after every member passes traversal/link/size checks."""
    destination.mkdir(parents=True, exist_ok=True)
    try:
        with ZipFile(archive_path) as archive:
            members = archive.infolist()
            if len(members) > MAX_ZIP_FILES:
                raise IntakeError(f"ZIP contains more than {MAX_ZIP_FILES} files", file=archive_path.name)
            if sum(member.file_size for member in members) > MAX_ZIP_UNCOMPRESSED_BYTES:
                raise IntakeError("ZIP uncompressed size exceeds the 50 MB safety limit", file=archive_path.name)
            for member in members:
                raw_name = member.filename
                posix = PurePosixPath(raw_name)
                windows = PureWindowsPath(raw_name)
                unix_mode = member.external_attr >> 16
                unsafe = (
                    not raw_name
                    or "\\" in raw_name
                    or posix.is_absolute()
                    or windows.is_absolute()
                    or windows.drive
                    or ".." in posix.parts
                    or stat.S_ISLNK(unix_mode)
                    or not _within(destination / Path(*posix.parts), destination)
                )
                if unsafe:
                    raise IntakeError(
                        f"unsafe ZIP member rejected: {raw_name!r}",
                        file=archive_path.name,
                        field="zip_member",
                    )
            corrupt = archive.testzip()
            if corrupt:
                raise IntakeError(f"ZIP member failed integrity check: {corrupt}", file=archive_path.name)
            for member in members:
                target = destination / Path(*PurePosixPath(member.filename).parts)
                if member.is_dir():
                    target.mkdir(parents=True, exist_ok=True)
                    continue
                target.parent.mkdir(parents=True, exist_ok=True)
                with archive.open(member) as source, target.open("wb") as output:
                    shutil.copyfileobj(source, output)
    except BadZipFile as exc:
        raise IntakeError("file is not a valid ZIP archive", file=archive_path.name) from exc


def _read_yaml(path: Path) -> dict:
    try:
        value = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    except (OSError, UnicodeError, yaml.YAMLError) as exc:
        raise IntakeError(f"invalid YAML: {exc}", file=path.name) from exc
    if not isinstance(value, dict):
        raise IntakeError("YAML document must be a mapping", file=path.name)
    return value


def _split_frontmatter(path: Path) -> tuple[dict, str]:
    text = path.read_text(encoding="utf-8")
    if not text.startswith("---"):
        raise IntakeError("Markdown must begin with YAML frontmatter", file=str(path), field="frontmatter")
    # Only horizontal whitespace belongs to the fence.  ``\s`` would consume
    # the blank line after it and therefore alter approved Markdown bodies.
    match = re.search(r"^---[ \t]*$", text[3:], flags=re.MULTILINE)
    if not match:
        raise IntakeError("Markdown frontmatter is not closed", file=str(path), field="frontmatter")
    end = 3 + match.end()
    try:
        meta = yaml.safe_load(text[3 : 3 + match.start()]) or {}
    except yaml.YAMLError as exc:
        raise IntakeError(f"invalid Markdown frontmatter: {exc}", file=str(path), field="frontmatter") from exc
    if not isinstance(meta, dict):
        raise IntakeError("Markdown frontmatter must be a mapping", file=str(path), field="frontmatter")
    return meta, text[end:]


def _render_frontmatter(meta: dict, untouched_body: str) -> str:
    return f"---\n{yaml.safe_dump(meta, sort_keys=False)}---{untouched_body}"


def _optional_frontmatter(path: Path) -> tuple[dict, str, bool]:
    """Return optional overview metadata without treating prose as routing."""
    text = path.read_text(encoding="utf-8")
    if not text.startswith("---"):
        return {}, text, False
    meta, body = _split_frontmatter(path)
    return meta, body, True


def _markdown_module_title(text: str) -> str | None:
    """Extract display-only title from a clearly labelled Markdown section."""
    match = re.search(
        r"^##[ \t]+Module title[ \t]*\n+(?:[ \t]*\n)*([^\n#].*?)\s*$",
        text,
        flags=re.IGNORECASE | re.MULTILINE,
    )
    if match:
        return match.group(1).strip()
    heading = re.search(r"^#[ \t]+(.+?)\s*$", text, flags=re.MULTILINE)
    return heading.group(1).strip() if heading else None


def _consensus(field: str, evidence: list[tuple[str, object]]) -> str:
    present = [(source, str(value).strip()) for source, value in evidence if str(value or "").strip()]
    values = {value for _source, value in present}
    if len(values) > 1:
        detail = ", ".join(f"{source}={value!r}" for source, value in present)
        raise IntakeError(f"conflicting structured {field} values: {detail}", field=field)
    if not values:
        raise IntakeError(f"insufficient structured metadata to determine {field}", field=field)
    return values.pop()


def _lesson_order_evidence(path: Path, meta: dict) -> list[tuple[str, int]]:
    evidence: list[tuple[str, int]] = []
    explicit = meta.get("lesson_order")
    if explicit not in (None, ""):
        try:
            evidence.append(("frontmatter lesson_order", int(explicit)))
        except (TypeError, ValueError) as exc:
            raise IntakeError(
                f"lesson_order must be a whole number, got {explicit!r}",
                file=path.name,
                field="lesson_order",
            ) from exc
    prefix = re.match(r"^(\d+)(?:[-_. ])", path.name)
    if prefix:
        evidence.append(("filename prefix", int(prefix.group(1))))
    lesson_number = meta.get("lesson_number")
    if lesson_number not in (None, ""):
        try:
            evidence.append(("frontmatter lesson_number", int(lesson_number)))
        except (TypeError, ValueError) as exc:
            raise IntakeError(
                f"lesson_number must be a whole number, got {lesson_number!r}",
                file=path.name,
                field="lesson_order",
            ) from exc
    title_number = re.match(r"^Lesson\s+(\d+)\b", str(meta.get("title") or ""), re.IGNORECASE)
    if title_number:
        evidence.append(("numbered lesson title", int(title_number.group(1))))
    return evidence


def _normalize_lesson_orders(raw_lessons: list[tuple[Path, dict, str]], notes: set[str]) -> None:
    orders: dict[int, str] = {}
    for path, meta, _body in raw_lessons:
        evidence = _lesson_order_evidence(path, meta)
        values = {value for _source, value in evidence}
        if not values:
            raise IntakeError(
                "lesson_order is missing and no numeric filename/title evidence is available",
                file=path.name,
                field="lesson_order",
            )
        if len(values) > 1:
            detail = ", ".join(f"{source}={value}" for source, value in evidence)
            raise IntakeError(
                f"conflicting lesson order evidence: {detail}",
                file=path.name,
                field="lesson_order",
            )
        order = values.pop()
        if order < 1:
            raise IntakeError("lesson_order must be at least 1", file=path.name, field="lesson_order")
        if order in orders:
            raise IntakeError(
                f"duplicate lesson_order {order} in {orders[order]} and {path.name}",
                file=path.name,
                field="lesson_order",
            )
        orders[order] = path.name
        if meta.get("lesson_order") in (None, ""):
            meta["lesson_order"] = order
            source = next(source for source, value in evidence if value == order)
            notes.add(f"{path.name}: lesson_order {order} derived from {source}")
    expected = list(range(1, len(raw_lessons) + 1))
    actual = sorted(orders)
    if actual != expected:
        raise IntakeError(
            f"lesson orders must form a contiguous 1..{len(raw_lessons)} sequence; got {actual}",
            field="lesson_order",
        )


def _tree_hash(path: Path) -> str:
    digest = hashlib.sha256()
    for item in sorted(path.rglob("*")):
        if item.is_symlink():
            raise IntakeError("symbolic links are not accepted in curriculum folders", file=str(item))
        if item.is_file():
            relative = item.relative_to(path).as_posix().encode()
            digest.update(len(relative).to_bytes(4, "big"))
            digest.update(relative)
            digest.update(item.read_bytes())
    return digest.hexdigest()


def _source_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest() if path.is_file() else _tree_hash(path)


def _slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def _normalize_scalar(value, aliases: dict[str, str], notes: set[str]):
    raw = str(value or "").strip()
    normalized = aliases.get(raw.casefold(), raw.lower().replace(" ", "_"))
    if raw and raw != normalized:
        notes.add(f"{raw} -> {normalized}")
    return normalized


def _json_cell(value) -> str:
    if value in (None, ""):
        return ""
    return json.dumps(value, ensure_ascii=False, separators=(",", ":")) if isinstance(value, (list, dict)) else str(value)


def _finding(
    category: str,
    message: str,
    *,
    file: str,
    field: str,
    row: int | None = None,
    key: str | None = None,
) -> dict:
    return {
        "category": category,
        "file": file,
        "row": row,
        "key": key,
        "field": field,
        "expected": None,
        "actual": None,
        "message": message,
    }


def _split_friendly_list(value) -> list[str]:
    if value in (None, ""):
        return []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    return [item.strip() for item in re.split(r"[;|]", str(value)) if item.strip()]


def _permission_status(value, notes: set[str]) -> str:
    raw = str(value or "unknown").strip()
    lowered = raw.casefold()
    if lowered in {"owned", "permitted", "requested", "unknown", "denied"}:
        return lowered
    if "nexus-authored" in lowered or "user-owned" in lowered:
        normalized = "owned"
    elif "permitted" in lowered or "official" in lowered or "public reference" in lowered:
        normalized = "permitted"
    else:
        normalized = "unknown"
    notes.add(f"{raw} -> {normalized}")
    return normalized


def _normalize_question_row(raw: dict, notes: set[str], *, quiz_title: str | None = None) -> dict:
    row: dict = {}
    for key, value in raw.items():
        source_key = str(key).strip()
        normalized_key = re.sub(r"[^a-z0-9]+", "_", source_key.casefold().replace("(s)", "s")).strip("_")
        canonical = FIELD_ALIASES.get(normalized_key, normalized_key)
        if canonical != source_key:
            notes.add(f"{source_key} -> {canonical}")
        row[canonical] = value
    row["question_type"] = _normalize_scalar(row.get("question_type"), QUESTION_TYPE_ALIASES, notes)
    row["importance"] = _normalize_scalar(row.get("importance"), IMPORTANCE_ALIASES, notes)
    row["permission_status"] = _permission_status(row.get("permission_status"), notes)
    editorial = str(row.get("final_validation_status") or "").strip()
    if editorial:
        normalized_editorial = EDITORIAL_STATUS_ALIASES.get(editorial.casefold(), editorial)
        if normalized_editorial != editorial:
            notes.add(f"editorial status {editorial} -> {normalized_editorial}")
        row["final_validation_status"] = normalized_editorial
    if quiz_title and not str(row.get("quiz_title") or "").strip():
        row["quiz_title"] = quiz_title
        notes.add("blank quiz_title -> module quiz blueprint title")

    tags = [tag.strip() for tag in str(row.get("tags") or "").split(",") if tag.strip()]
    for field in ("question_id", "lesson_id", "relationship", "style"):
        value = str(row.get(field) or "").strip()
        if value:
            tags.append(value)
    difficulty = str(row.get("difficulty") or "").strip()
    if difficulty and not difficulty.isdigit():
        tags.append(f"difficulty:{_slug(difficulty)}")
        row["difficulty"] = ""
        notes.add(f"difficulty label {difficulty!r} preserved as tag")
    row["tags"] = ",".join(dict.fromkeys(tags))

    if row["question_type"] == "short_answer":
        variants = _split_friendly_list(row.get("acceptable_answers"))
        if variants:
            row["acceptable_answers"] = variants
            notes.add("semicolon accepted_variants -> acceptable_answers list")
    if row["question_type"] == "free_response":
        concepts = _split_friendly_list(row.get("expected_concepts"))
        if concepts:
            row["expected_concepts"] = concepts
            notes.add("semicolon expected_concepts -> expected_concepts list")
        rubric = row.get("rubric")
        if isinstance(rubric, str) and rubric.strip() and not rubric.lstrip().startswith(("{", "[")):
            row["rubric"] = {"approved_text": rubric.strip()}
            notes.add("plain-text free-response rubric -> rubric.approved_text")
        minimum_raw = str(row.get("min_concepts_for_pass") or "").strip()
        minimum_match = re.match(r"^(\d+)\b", minimum_raw)
        if minimum_match:
            row["min_concepts_for_pass"] = int(minimum_match.group(1))
            if minimum_raw != minimum_match.group(1):
                notes.add(f"min_concepts_pass {minimum_raw!r} -> {minimum_match.group(1)}")
        guidance = str(row.get("partial_credit_guidance") or "").strip()
        if guidance:
            row["partial_credit"] = "true"
            row["rubric"] = {
                "partial_credit_guidance": guidance,
                "minimum_rule": minimum_raw,
            }
            notes.add("partial_credit_guidance -> free-response rubric")
    return row


def _normalize_quiz_blueprint(doc: dict, notes: set[str]) -> dict:
    """Translate author-facing quiz aliases into the existing V2 selector shape."""
    quiz = deepcopy(doc)
    for source, target in (
        ("title", "quiz_title"),
        ("question_count", "displayed_count"),
        ("pass_threshold_percent", "pass_percent"),
    ):
        if source in quiz and target not in quiz:
            quiz[target] = quiz[source]
            notes.add(f"module quiz {source} -> {target}")
    if quiz.get("pools") and not quiz.get("question_blueprint"):
        quiz["question_blueprint"] = [
            {
                "objective_codes": [str(pool["objective"])] if pool.get("objective") else [],
                "tags_any": [str(value) for value in pool.get("pool") or []],
                "count": int(pool.get("choose") or 0),
            }
            for pool in quiz["pools"]
        ]
        notes.add("module quiz pools -> question_blueprint ID-tag selectors")
    if quiz.get("required_category_coverage") and not quiz.get("category_requirements"):
        quiz["category_requirements"] = [
            {
                "category": str(row.get("category") or ""),
                "minimum": int(row.get("minimum") or 0),
                "tags_any": [str(value) for value in row.get("pool") or []],
            }
            for row in quiz["required_category_coverage"]
        ]
        notes.add("module quiz required_category_coverage -> category_requirements")
    return quiz


def _quick_check_rows(path: Path) -> list[dict]:
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return []
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Quick Checks" not in workbook.sheetnames:
        return []
    values = list(workbook["Quick Checks"].iter_rows(values_only=True))
    if not values:
        return []
    headers = []
    for value in values[0]:
        header = re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().casefold()).strip("_")
        headers.append({"quick_check_question_ids": "question_ids"}.get(header, header))
    return [dict(zip(headers, row)) for row in values[1:] if any(value is not None for value in row)]


def _attach_quick_checks(lessons: list[dict], rows: list[dict], notes: set[str]) -> None:
    """Attach workbook ID pools to lessons without inventing question content."""
    by_order = {int(item["lesson_order"]): item for item in lessons}
    by_title = {str(item["title"]).strip().casefold(): item for item in lessons}
    assigned: set[int] = set()
    for sheet_row, row in enumerate(rows, 2):
        match = re.fullmatch(r"L(\d+)", str(row.get("lesson_id") or "").strip(), re.IGNORECASE)
        order_value = row.get("lesson_order")
        order = int(match.group(1)) if match else int(order_value) if order_value not in (None, "") else None
        lesson = by_order.get(order) if order is not None else None
        lesson_key = str(row.get("lesson_key") or "").strip()
        key_lesson = next(
            (item for item in lessons if str(item.get("lesson_key") or "") == lesson_key),
            None,
        ) if lesson_key else None
        title = str(row.get("lesson_title") or "").strip()
        title_lesson = by_title.get(title.casefold()) if title else None
        candidates = [item for item in (lesson, key_lesson, title_lesson) if item is not None]
        if candidates and any(item is not candidates[0] for item in candidates[1:]):
            raise IntakeError(
                "Quick Check order, lesson key, and title identify different lessons",
                file="questions_and_editorial_review.xlsx",
                field=f"Quick Checks row {sheet_row}",
            )
        lesson = lesson or key_lesson or title_lesson
        if lesson is None:
            raise IntakeError(
                "Quick Check row does not resolve to a package lesson",
                file="questions_and_editorial_review.xlsx",
                field=f"Quick Checks row {sheet_row}",
            )
        lesson_order = int(lesson["lesson_order"])
        if lesson_order in assigned:
            raise IntakeError(
                f"duplicate Quick Check definition for lesson order {lesson_order}",
                file="questions_and_editorial_review.xlsx",
                field=f"Quick Checks row {sheet_row}",
            )
        ids = [value.strip() for value in str(row.get("question_ids") or "").split(",") if value.strip()]
        count = int(row.get("count") or len(ids))
        if count != len(ids):
            raise IntakeError(
                "Quick Check count does not match its question_ids",
                file="questions_and_editorial_review.xlsx",
                field=f"Quick Checks row {sheet_row}",
            )
        lesson["quick_check"] = {
            "title": f"Quick Check — {lesson['title']}",
            "displayed_count": count,
            "pass_percent": 60,
            "tags_any": ids,
        }
        assigned.add(lesson_order)
        notes.add(f"Quick Checks sheet row {sheet_row} -> lesson assessment metadata")


def _generated_service_desk_key(module_key: str, ordinal: int = 1) -> str:
    module_slug = _slug(module_key.removeprefix("module."))
    return f"curriculum-{module_slug}-service-desk-{ordinal:02d}"


def _has_inline_service_desk_content(doc: dict) -> bool:
    """Distinguish a stable-key reference from an approved inline definition."""
    return any(
        doc.get(field) not in (None, "", [], {})
        for field in (
            "ticket",
            "requester",
            "complaint",
            "business_impact",
            "stages",
            "grading_anchors",
            "correct_failure_behavior",
            "correct_outcomes",
        )
    )


def _service_desk_stage_map(stages, notes: set[str]) -> dict[str, list]:
    if isinstance(stages, dict):
        return deepcopy(stages)
    if not isinstance(stages, list):
        raise IntakeError(
            "inline Service Desk stages must be a mapping or ordered stage list",
            file="service_desk.yaml",
            field="stages",
        )
    normalized: dict[str, list] = {}
    for index, row in enumerate(stages, 1):
        if not isinstance(row, dict):
            raise IntakeError(
                f"inline Service Desk stage row {index} must be a mapping",
                file="service_desk.yaml",
                field="stages",
            )
        name = str(row.get("stage") or "").strip()
        expectations = row.get("expectations")
        if not name or not isinstance(expectations, list) or not expectations:
            raise IntakeError(
                f"inline Service Desk stage row {index} needs stage and expectations",
                file="service_desk.yaml",
                field="stages",
            )
        if name in normalized:
            raise IntakeError(
                f"duplicate inline Service Desk stage {name!r}",
                file="service_desk.yaml",
                field="stages",
            )
        normalized[name] = deepcopy(expectations)
    notes.add("Service Desk ordered stage list -> rubric dimension mapping")
    return normalized


def _normalize_inline_service_desk(
    doc: dict,
    module_key: str,
    valid_objectives: set[str],
    notes: set[str],
) -> tuple[str, dict]:
    """Map one approved curriculum scenario into the existing versioned engine."""
    approved_doc = deepcopy(doc)
    key = str(doc.get("scenario_key") or "").strip() or _generated_service_desk_key(module_key)
    if not doc.get("scenario_key"):
        notes.add(f"Service Desk scenario_key derived from module key -> {key}")
    for field in ("title", "stages", "grading_anchors"):
        if doc.get(field) in (None, "", [], {}):
            raise IntakeError(
                f"inline Service Desk field {field!r} is required",
                file="service_desk.yaml",
                field=field,
            )
    stage_map = _service_desk_stage_map(doc["stages"], notes)
    ticket = deepcopy(doc.get("ticket") or {})
    if not ticket:
        ticket = {
            "requester": deepcopy(doc.get("requester")),
            "complaint": doc.get("complaint"),
            "business_impact": doc.get("business_impact"),
            "initial_priority": doc.get("initial_priority") or doc.get("priority"),
            "twist": doc.get("twist"),
            "initial_facts": deepcopy(doc.get("initial_facts") or []),
        }
        notes.add("top-level Service Desk scenario fields -> existing engine ticket")
    for field in ("requester", "complaint", "business_impact"):
        value = ticket.get(field)
        if field == "requester" and isinstance(value, dict):
            value = value.get("name")
        if not str(ticket.get(field) or "").strip():
            if field == "requester" and value:
                continue
            raise IntakeError(
                f"inline Service Desk ticket field {field!r} is required",
                file="service_desk.yaml",
                field=field,
            )
    if not ticket.get("initial_priority"):
        ticket["initial_priority"] = "medium"
        notes.add("missing Service Desk priority -> existing engine default medium")
    required_stages = (
        "Investigation", "Diagnosis", "Remediation", "Verification", "Documentation"
    )
    if set(stage_map) != set(required_stages):
        raise IntakeError(
            f"inline Service Desk stages must be exactly {sorted(required_stages)}",
            file="service_desk.yaml",
            field="stages",
        )
    codes = [str(value) for value in doc.get("objectives") or []]
    for code in codes:
        if code not in valid_objectives:
            raise IntakeError(
                f"objective {code!r} is not defined for this certification version",
                file="service_desk.yaml",
                field="objectives",
            )
    ticket_id = key.upper()
    priority_text = str(ticket["initial_priority"])
    priority = next(
        (value for value in ("critical", "high", "medium", "low") if value in priority_text.casefold()),
        "medium",
    )
    anchors = deepcopy(doc["grading_anchors"])
    point_value = (
        sum(int(row.get("weight") or 0) for row in anchors)
        if isinstance(anchors, list)
        else 100
    ) or 100
    troubleshooting = [str(value) for value in ticket.get("initial_facts") or []]
    if ticket.get("twist"):
        troubleshooting.insert(0, str(ticket["twist"]))
    troubleshooting.extend(
        str(step)
        for stage in required_stages
        for step in stage_map.get(stage) or []
    )
    correct_failure_behavior = str(doc.get("correct_failure_behavior") or "").strip()
    correct_outcomes = [str(value) for value in doc.get("correct_outcomes") or []]
    if not correct_failure_behavior and not correct_outcomes:
        raise IntakeError(
            "inline Service Desk needs correct_failure_behavior or correct_outcomes",
            file="service_desk.yaml",
            field="correct_outcomes",
        )
    outcome_explanation = correct_failure_behavior or "\n".join(correct_outcomes)
    successful_outcomes = (
        ["pending", "escalated", "handed_off"]
        if correct_failure_behavior
        else correct_outcomes
    )
    placeholder = "Not specified by approved curriculum package"
    requester = ticket["requester"]
    requester_doc = requester if isinstance(requester, dict) else {}
    requester_name = str(requester_doc.get("name") or requester)
    hints = [
        {
            "id": f"hint-{index:02d}",
            "order": index,
            "pointPenalty": 0 if index == 1 else 5,
            "text": str(text),
        }
        for index, text in enumerate(doc.get("hints") or [], 1)
    ]
    if len(hints) < 3:
        raise IntakeError(
            "inline Service Desk scenario needs at least three approved hints",
            file="service_desk.yaml",
            field="hints",
        )
    definition = {
        "id": ticket_id,
        "title": doc["title"],
        "slug": key,
        "category": "service_desk",
        "priority": priority,
        "difficulty": "easy",
        "pointValue": point_value,
        "explanation": outcome_explanation,
        "description": {
            "issue": ticket["complaint"],
            "reportedByLine": requester_name,
            "businessImpact": ticket["business_impact"],
            "troubleshooting": troubleshooting,
        },
        "requester": {
            "name": requester_name,
            "department": str(requester_doc.get("department") or placeholder),
            "email": placeholder,
            "contact": placeholder,
            "location": str(requester_doc.get("location") or placeholder),
        },
        "device": {
            "assetTag": ticket_id,
            "deviceName": str(requester_doc.get("device") or placeholder),
            "kind": "mobile" if requester_doc.get("device") else "laptop",
            "operatingSystem": placeholder,
            "state": "active",
        },
        "sla": {"dueAt": placeholder, "target": priority_text},
        "initialWorldState": {
            "directoryOverlaySeeds": {},
            "assetOverlaySeeds": {},
            "chatMessageSeeds": [],
        },
        "objectives": [
            {
                "id": "professional-pending-or-completed-outcome",
                "description": outcome_explanation,
                "predicateType": "action_event_occurred",
                "predicateParams": {
                    "actionType": "ticket.add_note",
                    "payloadMatch": {"ticketId": ticket_id},
                },
                "required": True,
                "pointValue": point_value,
            }
        ],
        "requiredActions": [
            {
                "id": "document-professional-outcome",
                "actionType": "ticket.add_note",
                "description": outcome_explanation,
                "payloadMatch": {"ticketId": ticket_id},
            }
        ],
        "forbiddenActions": [],
        "hints": hints,
        "curriculum": approved_doc,
        "rubric_dimensions": stage_map,
        "grading_anchors": anchors,
        "successful_professional_outcomes": successful_outcomes,
    }
    return key, {
        "stable_key": key,
        "title": str(doc["title"]),
        "description": str(ticket["complaint"]),
        "category": "service_desk",
        "difficulty": 1,
        "definition": definition,
    }


class CurriculumIntakeProcessor:
    """Discover, validate, stage, archive, and integrate independent packages."""

    def __init__(
        self,
        *,
        dropbox_dir: Path,
        approved_dir: Path,
        content_dir: Path,
        backend_dir: Path,
        runtime_validator: Callable[[Path, dict], dict] | None = None,
        service_desk_keys: set[str] | None = None,
    ):
        self.dropbox_dir = Path(dropbox_dir)
        self.approved_dir = Path(approved_dir)
        self.content_dir = Path(content_dir)
        self.backend_dir = Path(backend_dir)
        self.runtime_validator = runtime_validator or self._validate_in_scratch_database
        self.service_desk_keys = service_desk_keys

    def discover(self) -> list[Path]:
        self.dropbox_dir.mkdir(parents=True, exist_ok=True)
        ignored = {"README.md", ".gitkeep", ".reports"}
        return sorted(
            [p for p in self.dropbox_dir.iterdir() if p.name not in ignored and (p.is_dir() or p.suffix.lower() == ".zip")],
            key=lambda p: p.name.casefold(),
        )

    def process(self, *, mode: str, allow_changed: bool = False) -> list[dict]:
        if mode not in {"check", "apply"}:
            raise ValueError("mode must be 'check' or 'apply'")
        results = []
        for source in self.discover():
            try:
                results.append(self._process_one(source, mode=mode, allow_changed=allow_changed))
            except Exception as exc:  # isolate one bad package from its siblings
                issue = exc.as_dict() if isinstance(exc, IntakeError) else IntakeError(str(exc)).as_dict()
                result = {
                    "source": source.name,
                    "status": "INVALID",
                    "errors": [issue],
                    "normalizations": [],
                    "package_hash": None,
                    "approved_destination": None,
                }
                self._write_report(source.name, result)
                results.append(result)
        return results

    def _process_one(self, source: Path, *, mode: str, allow_changed: bool) -> dict:
        if source.is_symlink():
            raise IntakeError("symbolic links are not accepted as inbox packages", file=source.name)
        package_hash = _source_hash(source)
        with tempfile.TemporaryDirectory(prefix="nexus-curriculum-intake-") as tmp:
            workspace = Path(tmp)
            if source.is_file():
                extracted = workspace / "extracted"
                safe_extract_zip(source, extracted)
                root = self._locate_package_root(extracted)
            else:
                root = self._locate_package_root(source)
            scan = self._scan_package(root, package_hash)
            if scan["errors"]:
                result = {
                    "source": source.name,
                    "status": "INVALID",
                    "errors": scan["errors"],
                    "warnings": scan["warnings"],
                    "findings": scan["findings"],
                    "normalizations": sorted(scan["normalizations"]),
                    "package_hash": package_hash,
                    "approved_destination": None,
                    "classification": scan.get("classification"),
                    "manifest": scan.get("manifest"),
                    "component_summary": scan["component_summary"],
                }
                self._write_report(source.name, result)
                return result
            normalized = self._normalize_and_validate(root, source.name, package_hash)
            module_root = (
                self.approved_dir
                / normalized["manifest"]["certification_key"]
                / normalized["manifest"]["version_key"]
                / normalized["manifest"]["module_key"]
            )
            classification = self._classify(module_root, package_hash)
            base = {
                "source": source.name,
                "errors": [],
                "warnings": scan["warnings"],
                "findings": scan["findings"],
                "normalizations": sorted(normalized["normalizations"]),
                "package_hash": package_hash,
                "approved_destination": str(module_root / package_hash),
                "classification": classification,
                "manifest": normalized["manifest"],
                "component_summary": scan["component_summary"],
            }
            staged_content = workspace / "content"
            shutil.copytree(self.content_dir, staged_content)
            runtime_destinations = self._write_runtime(normalized, staged_content)
            loader_result = self.runtime_validator(staged_content, normalized["manifest"])
            references = loader_result.get("references", {})
            unresolved = references.get("package_content_engine_unresolved")
            if unresolved is None:
                # Compatibility for injected/older validators that only
                # report the global loader result.
                unresolved = references.get("content_engine_unresolved", [])
            if unresolved:
                raise IntakeError(f"runtime references did not resolve: {unresolved}", field="assessments")
            service_unresolved = references.get("package_service_desk_unresolved", [])
            if service_unresolved:
                raise IntakeError(
                    f"Service Desk references did not resolve: {service_unresolved}",
                    field="service_desk",
                )

            if classification == "CHANGED" and not allow_changed:
                result = {**base, "status": "CHANGED_REQUIRES_REVIEW", "runtime_destinations": runtime_destinations}
                self._write_report(source.name, result)
                return result

            if mode == "check":
                if classification == "UNCHANGED":
                    status = "UNCHANGED"
                else:
                    status = "VALID_WITH_NORMALIZATION" if normalized["normalizations"] else "VALID"
                result = {**base, "status": status, "runtime_destinations": runtime_destinations}
                self._write_report(source.name, result)
                return result

            archive_destination = module_root / package_hash
            rollback = self._promote_runtime(staged_content, runtime_destinations)
            if classification == "UNCHANGED":
                try:
                    self._verify_archive(archive_destination, package_hash)
                    self._remove_inbox_item(source)
                except Exception:
                    self._rollback_runtime(rollback)
                    raise
                result = {**base, "status": "UNCHANGED", "runtime_destinations": runtime_destinations}
                self._write_report(source.name, result)
                return result
            try:
                receipt = self._archive(
                    source,
                    archive_destination,
                    normalized["manifest"],
                    loader_result,
                    runtime_destinations,
                    classification,
                )
                self._verify_archive(archive_destination, package_hash)
                if receipt["validation_result"] != "VALID":
                    raise IntakeError("receipt verification failed", field="receipt")
                self._remove_inbox_item(source)
            except Exception:
                self._rollback_runtime(rollback)
                if archive_destination.exists():
                    shutil.rmtree(archive_destination)
                raise
            result = {**base, "status": "IMPORTED", "runtime_destinations": runtime_destinations}
            self._write_report(source.name, result)
            return result

    @staticmethod
    def _locate_package_root(root: Path) -> Path:
        if (root / "module_overview.md").is_file():
            return root
        candidates = [p.parent for p in root.rglob("module_overview.md")]
        if len(candidates) != 1:
            raise IntakeError("package must contain exactly one module_overview.md", field="module_overview")
        return candidates[0]

    def _registry(self) -> dict:
        registry: dict[str, dict] = {}
        cert_dir = self.content_dir / "certifications"
        for path in sorted(cert_dir.glob("*.y*ml")):
            doc = _read_yaml(path)
            cert_key = str((doc.get("certification") or {}).get("cert_key") or "")
            if not cert_key:
                continue
            versions = {}
            for version in doc.get("versions") or []:
                version_key = str(version.get("version_key") or "")
                objective_file = self.content_dir / "objectives" / str(version.get("objectives_file") or "")
                objective_doc = _read_yaml(objective_file) if objective_file.is_file() else {}
                objectives = {str(row.get("code") or row.get("objective_code")) for row in objective_doc.get("objectives") or []}
                versions[version_key] = {
                    "data": version,
                    "objectives": objectives,
                    "domains": {str(row.get("domain_key")) for row in version.get("domains") or []},
                }
            registry[cert_key] = {"path": path, "doc": doc, "versions": versions}
        return registry

    def _scan_package(self, root: Path, package_hash: str) -> dict:
        """Collect independent compatibility findings without performing writes."""
        findings: list[dict] = []
        notes: set[str] = set()
        summary: dict = {
            "lessons": {},
            "questions": {},
            "resources": {},
            "quick_checks": {},
            "module_quiz": {},
            "practical": {},
            "service_desk": {},
            "explain": {},
            "provenance": {},
        }

        def add(category: str, message: str, *, file: str, field: str, row=None, key=None):
            findings.append(
                _finding(category, message, file=file, field=field, row=row, key=key)
            )

        overview: dict = {}
        overview_text = ""
        try:
            overview, overview_text, has_frontmatter = _optional_frontmatter(root / "module_overview.md")
            if not has_frontmatter:
                notes.add("module title derived from module_overview.md Markdown")
        except IntakeError as exc:
            add("BLOCKING_METADATA", str(exc), file="module_overview.md", field=exc.field)

        raw_lessons: list[tuple[Path, dict, str]] = []
        lesson_dir = root / "lessons"
        if not lesson_dir.is_dir():
            add("BLOCKING_METADATA", "lessons/ directory is required", file="lessons", field="lessons")
        else:
            for path in sorted(lesson_dir.glob("*.md")):
                try:
                    meta, body = _split_frontmatter(path)
                    raw_lessons.append((path, meta, body))
                except IntakeError as exc:
                    add("BLOCKING_METADATA", str(exc), file=path.name, field=exc.field)
        summary["lessons"]["count"] = len(raw_lessons)
        summary["lessons"]["keys"] = [
            str(meta.get("lesson_key") or "") for _path, meta, _body in raw_lessons
        ]
        summary["lessons"]["objectives"] = {
            path.name: [str(value) for value in meta.get("objectives") or []]
            for path, meta, _body in raw_lessons
        }
        summary["lessons"]["importance"] = {
            path.name: str(meta.get("importance") or "")
            for path, meta, _body in raw_lessons
        }
        summary["lessons"]["learning_relationships"] = {
            path.name: str(meta.get("learning_relationship") or "")
            for path, meta, _body in raw_lessons
        }

        # Ordering is independent per lesson; collect every contradiction and
        # then check the module-wide uniqueness/contiguity contract.
        order_rows: list[tuple[int, str]] = []
        for path, meta, _body in raw_lessons:
            try:
                evidence = _lesson_order_evidence(path, meta)
                values = {value for _source, value in evidence}
                if not values:
                    add(
                        "BLOCKING_METADATA",
                        "lesson_order is missing and no numeric filename/title evidence is available",
                        file=path.name,
                        field="lesson_order",
                    )
                    continue
                if len(values) > 1:
                    detail = ", ".join(f"{source}={value}" for source, value in evidence)
                    add(
                        "BLOCKING_METADATA",
                        f"conflicting lesson order evidence: {detail}",
                        file=path.name,
                        field="lesson_order",
                    )
                    continue
                order = values.pop()
                order_rows.append((order, path.name))
                if meta.get("lesson_order") in (None, ""):
                    source = next(source for source, value in evidence if value == order)
                    note = f"{path.name}: lesson_order {order} derived from {source}"
                    notes.add(note)
                    add("NORMALIZABLE", note, file=path.name, field="lesson_order")
            except IntakeError as exc:
                add("BLOCKING_METADATA", str(exc), file=path.name, field=exc.field)
        duplicates = sorted(order for order, _name in order_rows if sum(value == order for value, _ in order_rows) > 1)
        for order in sorted(set(duplicates)):
            names = [name for value, name in order_rows if value == order]
            add(
                "BLOCKING_METADATA",
                f"duplicate lesson_order {order}: {', '.join(names)}",
                file="lessons",
                field="lesson_order",
            )
        actual_orders = sorted({order for order, _name in order_rows})
        if order_rows and not duplicates and actual_orders != list(range(1, len(raw_lessons) + 1)):
            add(
                "BLOCKING_METADATA",
                f"lesson orders must form a contiguous 1..{len(raw_lessons)} sequence; got {actual_orders}",
                file="lessons",
                field="lesson_order",
            )
        summary["lessons"]["orders"] = [
            {"file": name, "order": order} for order, name in sorted(order_rows)
        ]

        provenance: dict = {}
        resources_doc: dict = {}
        for filename, target in (("provenance.yaml", "provenance"), ("resources.yaml", "resources")):
            path = root / filename
            if path.is_file():
                try:
                    doc = _read_yaml(path)
                    if target == "provenance":
                        provenance = doc
                    else:
                        resources_doc = doc
                except IntakeError as exc:
                    add("BLOCKING_METADATA", str(exc), file=filename, field=exc.field)

        version_key = module_key = domain = cert_key = None
        version_evidence = [("module_overview.md", overview.get("certification_version"))]
        version_evidence.extend((path.name, meta.get("certification_version")) for path, meta, _ in raw_lessons)
        version_evidence.extend(
            (f"resources.yaml:{row.get('resource_key') or index}", row.get("certification_version"))
            for index, row in enumerate(resources_doc.get("resources") or [], 1)
        )
        version_evidence.append(("provenance.yaml", provenance.get("certification_version")))
        module_evidence = [("module_overview.md", overview.get("module_key"))]
        module_evidence.extend((path.name, meta.get("module")) for path, meta, _ in raw_lessons)
        module_evidence.append(("provenance.yaml", provenance.get("module_key")))
        domain_evidence = [("module_overview.md", overview.get("domain"))]
        domain_evidence.extend((path.name, meta.get("domain")) for path, meta, _ in raw_lessons)
        for field, evidence in (
            ("certification_version", version_evidence),
            ("module_key", module_evidence),
            ("domain", domain_evidence),
        ):
            try:
                value = _consensus(field, evidence)
                if field == "certification_version":
                    version_key = value
                elif field == "module_key":
                    module_key = value
                else:
                    domain = value
            except IntakeError as exc:
                add("BLOCKING_METADATA", str(exc), file="package", field=field)

        registry = self._registry()
        version = None
        if version_key:
            owners = [cert for cert, item in registry.items() if version_key in item["versions"]]
            if len(owners) == 1:
                cert_key = owners[0]
                version = registry[cert_key]["versions"][version_key]
                explicit_cert = str(overview.get("certification") or "").strip()
                if explicit_cert and explicit_cert != cert_key:
                    add(
                        "BLOCKING_METADATA",
                        f"overview certification {explicit_cert!r} conflicts with hierarchy {cert_key!r}",
                        file="module_overview.md",
                        field="certification",
                    )
                elif not explicit_cert:
                    notes.add(f"certification resolved from {version_key} -> {cert_key}")
            else:
                add(
                    "BLOCKING_METADATA",
                    f"certification version {version_key!r} does not resolve uniquely in Nexus",
                    file="package",
                    field="certification_version",
                )
        if version and domain not in version["domains"]:
            add(
                "BLOCKING_METADATA",
                f"domain {domain!r} does not belong to {version_key!r}",
                file="package",
                field="domain",
            )
        title = str(overview.get("title") or "").strip() or _markdown_module_title(overview_text)
        if not title:
            add("BLOCKING_METADATA", "module title could not be resolved", file="module_overview.md", field="title")

        valid_objectives = version["objectives"] if version else set()
        objective_codes: set[str] = set()
        lesson_keys = {str(meta.get("lesson_key")) for _path, meta, _body in raw_lessons if meta.get("lesson_key")}
        for path, meta, _body in raw_lessons:
            for field in ("lesson_key", "title", "importance", "learning_relationship", "objectives"):
                if meta.get(field) in (None, "", []):
                    add("BLOCKING_METADATA", f"lesson field {field!r} is required", file=path.name, field=field)
            for code in [str(value) for value in meta.get("objectives") or []]:
                objective_codes.add(code)
                if valid_objectives and code not in valid_objectives:
                    add("BLOCKING_CONTENT", f"objective {code!r} is not defined for {version_key}", file=path.name, field="objective")

        # Question workbook and its supporting sheets are inspected even when
        # lesson metadata has unrelated errors.
        raw_questions: list[dict] = []
        question_ids: set[str] = set()
        question_types: dict[str, int] = {}
        question_objectives: dict[str, int] = {}
        question_importance: dict[str, int] = {}
        question_provenance: dict[str, int] = {}
        short_answer_count = 0
        free_response_count = 0
        question_items: list[dict] = []
        normalized_editorial_statuses: set[str] = set()
        workbook = None
        try:
            question_path = self._question_path(root)
            raw_questions = (
                parse_xlsx_file(question_path.read_bytes())
                if question_path.suffix.lower() in {".xlsx", ".xlsm"}
                else parse_csv_file(question_path.read_bytes())
            )
            if question_path.suffix.lower() in {".xlsx", ".xlsm"}:
                workbook = openpyxl.load_workbook(question_path, read_only=True, data_only=True)
        except Exception as exc:
            add("BLOCKING_METADATA", f"question workbook could not be read: {exc}", file="questions", field="questions")
            question_path = root / "questions"
        quiz_doc = {}
        try:
            quiz_doc = _read_yaml(root / "module_quiz_blueprint.yaml")
        except IntakeError as exc:
            add("BLOCKING_METADATA", str(exc), file="module_quiz_blueprint.yaml", field=exc.field)
        quiz_title = str(quiz_doc.get("quiz_title") or quiz_doc.get("title") or "").strip() or None
        for row_number, raw in enumerate(raw_questions, 2):
            row_notes: set[str] = set()
            row = _normalize_question_row(raw, row_notes, quiz_title=quiz_title)
            notes.update(row_notes)
            if row.get("final_validation_status"):
                normalized_editorial_statuses.add(
                    str(row["final_validation_status"]).strip()
                )
            qid = str(row.get("question_id") or "").strip()
            if qid:
                if qid in question_ids:
                    add("BLOCKING_METADATA", f"duplicate question_id {qid}", file=question_path.name, field="question_id", row=row_number, key=qid)
                question_ids.add(qid)
            qtype = str(row.get("question_type") or "")
            question_types[qtype] = question_types.get(qtype, 0) + 1
            codes = [part.strip() for part in str(row.get("objective_code") or "").split(",") if part.strip()]
            objective_codes.update(codes)
            for code in codes:
                question_objectives[code] = question_objectives.get(code, 0) + 1
            importance = str(row.get("importance") or "")
            question_importance[importance] = question_importance.get(importance, 0) + 1
            source = str(row.get("source_name") or "")
            question_provenance[source] = question_provenance.get(source, 0) + 1
            if qtype == "short_answer" and row.get("acceptable_answers"):
                short_answer_count += 1
            if qtype == "free_response" and all(
                row.get(field)
                for field in ("expected_concepts", "min_concepts_for_pass", "rubric_version")
            ):
                free_response_count += 1
            for code in codes:
                if valid_objectives and code not in valid_objectives:
                    add("BLOCKING_CONTENT", f"objective {code!r} is not defined for {version_key}", file=question_path.name, field="objective", row=row_number, key=qid or None)
            question_items.append(
                {
                    "id": qid,
                    "objective_codes": codes,
                    "tags": [tag for tag in str(row.get("tags") or "").split(",") if tag],
                }
            )
            row.update(
                certification=cert_key,
                certification_version=version_key,
                domain=domain,
                module=module_key,
            )
            validation = validate_question(row)
            for issue in validation.errors:
                add("BLOCKING_CONTENT", issue.message, file=question_path.name, field=issue.field, row=row_number, key=qid or None)
            if not str(row.get("explanation") or "").strip():
                add("BLOCKING_CONTENT", "question explanation is required", file=question_path.name, field="explanation", row=row_number, key=qid or None)
            difficulty = str(raw.get("difficulty") or "").strip()
            if difficulty and not difficulty.isdigit():
                add(
                    "WARNING",
                    f"friendly difficulty {difficulty!r} has no canonical 1-5 mapping and will be retained as a tag",
                    file=question_path.name,
                    field="difficulty",
                    row=row_number,
                    key=qid or None,
                )
        summary["questions"] = {
            "count": len(raw_questions),
            "type_distribution": question_types,
            "objective_distribution": question_objectives,
            "importance_distribution": question_importance,
            "provenance_distribution": question_provenance,
            "short_answer_rows_with_variants": short_answer_count,
            "free_response_rows_with_rubric": free_response_count,
            "answers_and_explanations_validated": not any(
                row["file"] == question_path.name
                and row["category"] == "BLOCKING_CONTENT"
                and row["field"] in {"correct_answers", "explanation"}
                for row in findings
            ),
            "ids": sorted(question_ids),
        }

        # Quick Checks live in an optional workbook sheet in author-friendly packages.
        quick_rows: list[dict] = []
        if workbook and "Quick Checks" in workbook.sheetnames:
            values = list(workbook["Quick Checks"].iter_rows(values_only=True))
            if values:
                headers = []
                for value in values[0]:
                    header = re.sub(
                        r"[^a-z0-9]+",
                        "_",
                        str(value or "").strip().casefold(),
                    ).strip("_")
                    headers.append(
                        {"quick_check_question_ids": "question_ids"}.get(header, header)
                    )
                quick_rows = [dict(zip(headers, row)) for row in values[1:] if any(value is not None for value in row)]
            for index, row in enumerate(quick_rows, 2):
                ids = [value.strip() for value in str(row.get("question_ids") or "").split(",") if value.strip()]
                missing = sorted(set(ids) - question_ids)
                if missing:
                    add("BLOCKING_METADATA", f"Quick Check references unknown question IDs: {missing}", file=question_path.name, field="question_ids", row=index)
                if int(row.get("count") or 0) != len(ids):
                    add("BLOCKING_METADATA", "Quick Check count does not match its question_ids", file=question_path.name, field="count", row=index)
                notes.add(f"Quick Checks sheet row {index} -> lesson assessment metadata")
        if workbook:
            workbook.close()
        summary["quick_checks"] = {"count": len(quick_rows), "rows": quick_rows}

        # Resources use the canonical loader shape already; inspect all links.
        resource_rows = resources_doc.get("resources") or []
        for index, row in enumerate(resource_rows, 1):
            key = str(row.get("resource_key") or "").strip()
            resource_type = row.get("resource_type") or row.get("type")
            permission = str(row.get("permission_status") or "").strip()
            if permission:
                _permission_status(permission, notes)
            if row.get("type") and not row.get("resource_type"):
                note = f"resource {key or index}: type -> resource_type"
                notes.add(note)
                add("NORMALIZABLE", note, file="resources.yaml", field="resource_type", row=index, key=key or None)
            for field in ("resource_key", "title", "resource_type", "url", "provider"):
                value = resource_type if field == "resource_type" else row.get(field)
                if not str(value or "").strip():
                    add("BLOCKING_METADATA", f"resource field {field!r} is required", file="resources.yaml", field=field, row=index, key=key or None)
            for link in row.get("links") or []:
                lesson_key = str(link.get("lesson_key") or "")
                if lesson_key and lesson_key not in lesson_keys:
                    add("BLOCKING_METADATA", f"resource references unknown lesson {lesson_key!r}", file="resources.yaml", field="lesson_key", row=index, key=key or None)
                linked_module = str(link.get("module_key") or "")
                if linked_module and module_key and linked_module != module_key:
                    add("BLOCKING_METADATA", f"resource module {linked_module!r} conflicts with {module_key!r}", file="resources.yaml", field="module_key", row=index, key=key or None)
        summary["resources"] = {
            "count": len(resource_rows),
            "links": sum(len(row.get("links") or []) for row in resource_rows),
            "lesson_mappings": {
                str(row.get("resource_key") or ""): [
                    str(link.get("lesson_key") or "") for link in row.get("links") or []
                ]
                for row in resource_rows
            },
        }

        # Explain prompts can be normalized mechanically except objective links,
        # which must be explicit and are never inferred from prose.
        prompt_rows: list[dict] = []
        try:
            prompt_doc = _read_yaml(root / "explain_prompts.yaml") if (root / "explain_prompts.yaml").is_file() else {}
            prompt_rows = prompt_doc.get("prompts") or []
            top_rubric_version = prompt_doc.get("rubric_version")
            for index, row in enumerate(prompt_rows, 1):
                prompt_id = str(row.get("prompt_key") or row.get("id") or "").strip()
                if not row.get("prompt_key") and row.get("id") and module_key:
                    notes.add(f"Explain {row['id']}: stable prompt_key generated from module key and id")
                if not row.get("rubric_version") and top_rubric_version:
                    notes.add(f"Explain {prompt_id}: inherited package rubric_version")
                if not row.get("rubric") and row.get("minimum_for_pass") is not None:
                    notes.add(f"Explain {prompt_id}: minimum_for_pass -> rubric metadata")
                codes = [str(code) for code in row.get("objectives") or []]
                if not codes:
                    add("BLOCKING_METADATA", "Explain prompt needs explicit objective relationships", file="explain_prompts.yaml", field="objectives", row=index, key=prompt_id or None)
                for code in codes:
                    if valid_objectives and code not in valid_objectives:
                        add("BLOCKING_CONTENT", f"objective {code!r} is not defined for {version_key}", file="explain_prompts.yaml", field="objectives", row=index, key=prompt_id or None)
        except IntakeError as exc:
            add("BLOCKING_METADATA", str(exc), file="explain_prompts.yaml", field=exc.field)
        summary["explain"] = {
            "count": len(prompt_rows),
            "prompt_keys": [
                str(row.get("prompt_key") or row.get("id") or "") for row in prompt_rows
            ],
            "objectives": {
                str(row.get("prompt_key") or row.get("id") or ""): [
                    str(code) for code in row.get("objectives") or []
                ]
                for row in prompt_rows
            },
            "concept_and_rubric_rows": sum(
                bool(row.get("expected_concepts"))
                and bool(row.get("rubric") or row.get("minimum_for_pass") is not None)
                for row in prompt_rows
            ),
        }

        # Practical top-level documents have a direct LabTemplate mapping.
        practical_doc: dict = {}
        if (root / "practical.yaml").is_file():
            try:
                practical_doc = _read_yaml(root / "practical.yaml")
                if practical_doc.get("title") and not practical_doc.get("practical") and not practical_doc.get("labs"):
                    notes.add("top-level practical fields -> LabTemplate")
                for code in [str(value) for value in practical_doc.get("objectives") or []]:
                    if valid_objectives and code not in valid_objectives:
                        add("BLOCKING_CONTENT", f"objective {code!r} is not defined for {version_key}", file="practical.yaml", field="objectives")
            except IntakeError as exc:
                add("BLOCKING_METADATA", str(exc), file="practical.yaml", field=exc.field)
        summary["practical"] = {
            "present": bool(practical_doc),
            "title": practical_doc.get("title")
            or (practical_doc.get("practical") or {}).get("title"),
            "objectives": [str(value) for value in practical_doc.get("objectives") or []],
            "mapping": "LabTemplate" if practical_doc else None,
            "vm_required": bool(practical_doc.get("proxmox_template_vmid")),
        }

        # Nexus has one Service Desk engine. Inline definitions cannot silently
        # become a different scenario or reuse a key with different outcomes.
        service_doc: dict = {}
        service_stage_names: list[str] = []
        if (root / "service_desk.yaml").is_file():
            try:
                service_doc = _read_yaml(root / "service_desk.yaml")
                stable_key = str(service_doc.get("scenario_key") or "").strip()
                if _has_inline_service_desk_content(service_doc):
                    try:
                        stable_key, scenario = _normalize_inline_service_desk(
                            service_doc,
                            str(module_key or ""),
                            valid_objectives,
                            notes,
                        )
                        service_stage_names = list(
                            (scenario["definition"].get("rubric_dimensions") or {}).keys()
                        )
                        add(
                            "NORMALIZABLE",
                            f"inline Service Desk scenario -> existing engine key {stable_key}",
                            file="service_desk.yaml",
                            field="scenario_key",
                            key=stable_key,
                        )
                    except IntakeError as exc:
                        add(
                            "BLOCKING_METADATA",
                            str(exc),
                            file="service_desk.yaml",
                            field=exc.field,
                        )
                elif self.service_desk_keys is not None and stable_key not in self.service_desk_keys:
                    add("BLOCKING_METADATA", f"Service Desk scenario {stable_key!r} is not registered", file="service_desk.yaml", field="scenario_key")
                for code in [
                    str(value)
                    for field in ("objectives", "reinforces")
                    for value in service_doc.get(field) or []
                ]:
                    if valid_objectives and code not in valid_objectives:
                        add("BLOCKING_CONTENT", f"objective {code!r} is not defined for {version_key}", file="service_desk.yaml", field="objectives")
            except IntakeError as exc:
                add("BLOCKING_METADATA", str(exc), file="service_desk.yaml", field=exc.field)
        summary["service_desk"] = {
            "present": bool(service_doc),
            "scenario_key": stable_key if service_doc else None,
            "objectives": [str(value) for value in service_doc.get("objectives") or []],
            "reinforces": [str(value) for value in service_doc.get("reinforces") or []],
            "rubric_dimensions": service_stage_names,
        }

        # Human quiz aliases and ID pools map to current blueprint selectors.
        if quiz_doc:
            if "title" in quiz_doc and "quiz_title" not in quiz_doc:
                notes.add("module quiz title -> quiz_title")
            if "question_count" in quiz_doc and "displayed_count" not in quiz_doc:
                notes.add("module quiz question_count -> displayed_count")
            if "pass_threshold_percent" in quiz_doc and "pass_percent" not in quiz_doc:
                notes.add("module quiz pass_threshold_percent -> pass_percent")
            pools = quiz_doc.get("pools") or []
            for index, pool in enumerate(pools, 1):
                ids = {str(value) for value in pool.get("pool") or []}
                missing = sorted(ids - question_ids)
                if missing:
                    add("BLOCKING_METADATA", f"quiz pool references unknown question IDs: {missing}", file="module_quiz_blueprint.yaml", field="pool", row=index)
                code = str(pool.get("objective") or "")
                if valid_objectives and code and code not in valid_objectives:
                    add("BLOCKING_CONTENT", f"objective {code!r} is not defined for {version_key}", file="module_quiz_blueprint.yaml", field="objective", row=index)
            if pools:
                notes.add("module quiz pools -> question_blueprint ID-tag selectors")
            normalized_quiz = _normalize_quiz_blueprint(quiz_doc, notes)
            categories = normalized_quiz.get("category_requirements") or []
            if categories:
                try:
                    selected = select_constrained(
                        question_items,
                        normalized_quiz.get("question_blueprint") or [],
                        categories,
                        excluded_ids={
                            str(value)
                            for value in normalized_quiz.get("exclude_question_ids") or []
                        },
                        rng=random.Random(0),
                    )
                    displayed = int(normalized_quiz.get("displayed_count") or 0)
                    if len(selected) != displayed:
                        raise ConstraintSelectionError(
                            "objective quota total does not match displayed_count"
                        )
                except ConstraintSelectionError as exc:
                    add(
                        "BLOCKING_METADATA",
                        f"Module Quiz constraints are not satisfiable: {exc}",
                        file="module_quiz_blueprint.yaml",
                        field="required_category_coverage",
                    )
        summary["module_quiz"] = {
            "title": quiz_doc.get("quiz_title") or quiz_doc.get("title"),
            "displayed_count": quiz_doc.get("displayed_count") or quiz_doc.get("question_count"),
            "pass_percent": quiz_doc.get("pass_percent") or quiz_doc.get("pass_threshold_percent"),
            "pool_count": len(quiz_doc.get("pools") or quiz_doc.get("question_blueprint") or []),
            "objective_distribution": quiz_doc.get("objective_distribution") or {},
            "required_category_coverage": quiz_doc.get("required_category_coverage") or [],
        }

        # Editorial status may be proven by unanimous row-level review records.
        editorial = str(provenance.get("editorial_status") or overview.get("editorial_status") or "").strip().lower()
        statuses = normalized_editorial_statuses
        approved_statuses = {"APPROVED", "APPROVED_AFTER_EDIT"}
        if editorial != "validated":
            if statuses and statuses.issubset(approved_statuses):
                notes.add("unanimous workbook final_validation_status -> editorial_status validated")
            else:
                add("BLOCKING_CONTENT", "editorial validation is not established", file="provenance.yaml", field="editorial_status")
        reviewed = (provenance.get("question_bank_summary") or {}).get("total_questions") or provenance.get("reviewed_question_count")
        if reviewed is not None and int(reviewed) != len(raw_questions):
            add("BLOCKING_CONTENT", f"reviewed count {reviewed} does not match {len(raw_questions)} question rows", file="provenance.yaml", field="reviewed_question_count")
        summary["provenance"] = {
            "editorial_status": "validated" if editorial == "validated" or (statuses and statuses.issubset(approved_statuses)) else editorial,
            "reviewed_question_count": int(reviewed or len(raw_questions)),
        }

        for required_file in ("CONTENT_STATUS.md", "question_bank_quality_rules.yaml"):
            if not (root / required_file).is_file():
                add("WARNING", f"optional review artifact {required_file} is absent", file=required_file, field="file")

        # Every automatic translation appears in both the concise normalization
        # list and the categorized findings. This makes check-mode output useful
        # without forcing callers to reconcile two reporting models.
        reported_normalizations = {
            row["message"] for row in findings if row["category"] == "NORMALIZABLE"
        }
        for note in sorted(notes - reported_normalizations):
            add("NORMALIZABLE", note, file="package", field="normalization")

        manifest = {
            "source_package_hash": package_hash,
            "certification_key": cert_key,
            "version_key": version_key,
            "domain_key": domain,
            "module_key": module_key,
            "module_title": title,
            "objectives": sorted(objective_codes),
            "lesson_count": len(raw_lessons),
            "question_count": len(raw_questions),
            "resource_count": len(resource_rows),
            "explain_count": len(prompt_rows),
            "practical_present": bool(practical_doc),
            "service_desk_present": bool(service_doc),
        }
        classification = None
        if cert_key and version_key and module_key:
            classification = self._classify(
                self.approved_dir / cert_key / version_key / module_key,
                package_hash,
            )
        routing_priority = {
            "certification_version": 0,
            "certification": 1,
            "domain": 2,
            "module_key": 3,
        }
        errors = sorted(
            (row for row in findings if row["category"].startswith("BLOCKING_")),
            key=lambda row: routing_priority.get(row["field"], 10),
        )
        warnings = [row for row in findings if row["category"] == "WARNING"]
        return {
            "findings": findings,
            "errors": errors,
            "warnings": warnings,
            "normalizations": notes,
            "manifest": manifest,
            "classification": classification,
            "component_summary": summary,
        }

    def _normalize_and_validate(self, root: Path, original_name: str, package_hash: str) -> dict:
        notes: set[str] = set()
        overview_path = root / "module_overview.md"
        overview, overview_text, has_overview_frontmatter = _optional_frontmatter(overview_path)

        lesson_dir = root / "lessons"
        if not lesson_dir.is_dir():
            raise IntakeError("lessons/ directory is required", field="lessons")
        raw_lessons: list[tuple[Path, dict, str]] = []
        for path in sorted(lesson_dir.glob("*.md")):
            meta, body = _split_frontmatter(path)
            raw_lessons.append((path, meta, body))
        if not raw_lessons:
            raise IntakeError("at least one lesson Markdown file is required", field="lessons")
        _normalize_lesson_orders(raw_lessons, notes)

        provenance = _read_yaml(root / "provenance.yaml") if (root / "provenance.yaml").is_file() else {}
        resources_doc = _read_yaml(root / "resources.yaml") if (root / "resources.yaml").is_file() else {}
        version_evidence = [("module_overview.md", overview.get("certification_version"))]
        version_evidence.extend((path.name, meta.get("certification_version")) for path, meta, _body in raw_lessons)
        version_evidence.extend(
            (f"resources.yaml:{row.get('resource_key') or index}", row.get("certification_version"))
            for index, row in enumerate(resources_doc.get("resources") or [], 1)
        )
        version_evidence.append(("provenance.yaml", provenance.get("certification_version")))
        module_evidence = [("module_overview.md", overview.get("module_key"))]
        module_evidence.extend((path.name, meta.get("module")) for path, meta, _body in raw_lessons)
        module_evidence.append(("provenance.yaml", provenance.get("module_key")))
        domain_evidence = [("module_overview.md", overview.get("domain"))]
        domain_evidence.extend((path.name, meta.get("domain")) for path, meta, _body in raw_lessons)

        version_key = _consensus("certification_version", version_evidence)
        module_key = _consensus("module_key", module_evidence)
        domain = _consensus("domain", domain_evidence)
        registry = self._registry()
        owners = [cert for cert, item in registry.items() if version_key in item["versions"]]
        if len(owners) != 1:
            message = (
                f"version '{version_key}' is not in Nexus certification data"
                if not owners
                else f"version '{version_key}' belongs to multiple certifications: {owners}"
            )
            raise IntakeError(message, field="certification_version")
        cert_key = owners[0]
        explicit_cert = str(overview.get("certification") or "").strip()
        if explicit_cert and explicit_cert != cert_key:
            raise IntakeError(
                f"conflicting structured certification values: module_overview.md={explicit_cert!r}, hierarchy={cert_key!r}",
                field="certification",
            )
        if not explicit_cert:
            notes.add(f"certification resolved from {version_key} -> {cert_key}")
        version = registry[cert_key]["versions"][version_key]
        if domain not in version["domains"]:
            raise IntakeError(f"domain '{domain}' does not belong to '{version_key}'", field="domain")

        structured_title = str(overview.get("title") or "").strip()
        prose_title = _markdown_module_title(overview_text)
        module_title = structured_title or prose_title
        if not module_title:
            raise IntakeError("module title is unavailable in structured metadata or the labelled Markdown overview", file="module_overview.md", field="title")
        if not has_overview_frontmatter:
            notes.add("module title derived from module_overview.md Markdown")

        lessons: list[tuple[str, str]] = []
        lesson_keys: set[str] = set()
        objective_codes: set[str] = set()
        lesson_meta_rows = []
        for path, meta, body in raw_lessons:
            for field in ("lesson_key", "title", "importance", "learning_relationship", "objectives"):
                if meta.get(field) in (None, "", []):
                    raise IntakeError(f"lesson field '{field}' is required", file=path.name, field=field)
            if str(meta.get("certification_version")) != version_key or str(meta.get("module")) != module_key:
                raise IntakeError("lesson version/module does not match package metadata", file=path.name, field="module")
            meta["domain"] = str(meta.get("domain") or domain)
            meta["importance"] = _normalize_scalar(meta["importance"], IMPORTANCE_ALIASES, notes)
            meta["learning_relationship"] = _normalize_scalar(meta["learning_relationship"], RELATIONSHIP_ALIASES, notes)
            lesson_objectives = [str(value) for value in meta.get("objectives") or []]
            self._validate_objectives(lesson_objectives, version["objectives"], path.name)
            objective_codes.update(lesson_objectives)
            key = str(meta["lesson_key"])
            if key in lesson_keys:
                raise IntakeError(f"duplicate lesson key '{key}'", file=path.name, field="lesson_key")
            lesson_keys.add(key)
            lessons.append((path.name, _render_frontmatter(meta, body)))
            lesson_meta_rows.append(meta)
        quiz = _normalize_quiz_blueprint(
            _read_yaml(root / "module_quiz_blueprint.yaml"), notes
        )
        quiz_title = str(quiz.get("quiz_title") or "").strip()
        question_path = self._question_path(root)
        _attach_quick_checks(lesson_meta_rows, _quick_check_rows(question_path), notes)
        # Re-render after workbook Quick Checks have been attached to normalized
        # lesson frontmatter. Approved Markdown on disk remains untouched.
        lessons = [
            (path.name, _render_frontmatter(meta, body))
            for path, meta, body in raw_lessons
        ]
        raw_rows = parse_xlsx_file(question_path.read_bytes()) if question_path.suffix.lower() in {".xlsx", ".xlsm"} else parse_csv_file(question_path.read_bytes())
        normalized_rows = []
        for row_number, raw in enumerate(raw_rows, 2):
            row = _normalize_question_row(raw, notes, quiz_title=quiz_title)
            row["certification"] = cert_key
            row["certification_version"] = version_key
            row["domain"] = domain
            row["module"] = module_key
            if row.get("published") in (None, ""):
                row["published"] = "true"
                notes.add("blank published -> true for validated package")
            payload = row_to_payload(row)
            validation = validate_question(payload)
            if not validation.valid:
                message = "; ".join(issue.message for issue in validation.errors)
                raise IntakeError(message, file=question_path.name, field=f"row {row_number}")
            row_objectives = parse_objective_codes(row.get("objective_code"))
            self._validate_objectives(row_objectives, version["objectives"], question_path.name)
            objective_codes.update(row_objectives)
            normalized_rows.append({key: _json_cell(row.get(key)) for key in TEMPLATE_COLUMNS})
        if not normalized_rows:
            raise IntakeError("question bank is empty", file=question_path.name, field="questions")

        editorial = str(provenance.get("editorial_status") or overview.get("editorial_status") or "").lower()
        statuses = {
            str(_normalize_question_row(row, set(), quiz_title=quiz_title).get("final_validation_status") or "").strip()
            for row in raw_rows
        }
        if editorial != "validated" and statuses and statuses.issubset(
            {"APPROVED", "APPROVED_AFTER_EDIT"}
        ):
            editorial = "validated"
            notes.add("unanimous workbook final_validation_status -> editorial_status validated")
        if editorial != "validated":
            raise IntakeError("package must have validated editorial status", file="provenance.yaml", field="editorial_status")
        reviewed_count = int(
            (provenance.get("question_bank_summary") or {}).get("total_questions")
            or provenance.get("reviewed_question_count")
            or len(normalized_rows)
        )
        if reviewed_count != len(normalized_rows):
            raise IntakeError("reviewed question count does not match workbook rows", file="provenance.yaml", field="reviewed_question_count")

        quiz_title = str(quiz.get("quiz_title") or normalized_rows[0].get("quiz_title") or "").strip()
        if not quiz_title:
            raise IntakeError("module quiz title is required", file="module_quiz_blueprint.yaml", field="quiz_title")
        for entry in quiz.get("question_blueprint") or []:
            self._validate_objectives([str(x) for x in entry.get("objective_codes") or []], version["objectives"], "module_quiz_blueprint.yaml")
        self._validate_assessment_pools(lesson_meta_rows, quiz, normalized_rows)

        resources = self._normalize_resources(root, version_key, lesson_keys, notes)
        prompts = self._normalize_prompts(root, version_key, module_key, domain, version["objectives"], notes)
        labs, practical_title = self._normalize_practical(root, notes)
        service_desk_key, service_desk_scenarios = self._service_desk(
            root, module_key, version["objectives"], notes
        )
        assessments = self._build_assessments(module_key, lesson_meta_rows, quiz, quiz_title, practical_title, service_desk_key, bool(prompts))
        existing_module = next(
            (
                row
                for row in version["data"].get("modules") or []
                if str(row.get("module_key") or "") == module_key
            ),
            None,
        )
        module = {
            "module_key": module_key,
            "title": module_title,
            "skill_promise": str(overview.get("skill_promise") or "") or None,
            "certification_domain_key": domain,
            "importance_hint": _normalize_scalar(overview.get("importance", "working_knowledge"), IMPORTANCE_ALIASES, notes),
            "display_order": int(
                overview.get("display_order")
                or (existing_module or {}).get("display_order")
                or max((int(row.get("display_order") or 0) for row in version["data"].get("modules") or []), default=0) + 1
            ),
            "assessments": assessments,
        }
        manifest = {
            "manifest_version": 1,
            "source_package_hash": package_hash,
            "original_filename": original_name,
            "certification_key": cert_key,
            "version_key": version_key,
            "domain_key": domain,
            "module_key": module_key,
            "module_title": module["title"],
            "lessons": sorted(lesson_keys),
            "objectives": sorted(objective_codes),
            "lesson_count": len(lessons),
            "question_count": len(normalized_rows),
            "resource_count": len(resources),
            "explain_count": len(prompts),
            "practical_present": bool(labs),
            "service_desk_present": bool(service_desk_key),
            "service_desk_keys": [service_desk_key] if service_desk_key else [],
            "provenance": provenance,
        }
        return {
            "manifest": manifest,
            "normalizations": notes,
            "registry_path": registry[cert_key]["path"].name,
            "module": module,
            "lessons": lessons,
            "questions": normalized_rows,
            "quiz_title": quiz_title,
            "resources": resources,
            "prompts": prompts,
            "labs": labs,
            "service_desk_scenarios": service_desk_scenarios,
        }

    @staticmethod
    def _validate_objectives(codes: list[str], valid: set[str], filename: str) -> None:
        for code in codes:
            if not code or code not in valid:
                raise IntakeError(f"objective '{code}' is not defined for this certification version", file=filename, field="objective")

    @staticmethod
    def _question_path(root: Path) -> Path:
        candidates = [
            root / "questions_and_editorial_review.xlsx",
            root / "questions_and_editorial_review.csv",
            root / "questions.xlsx",
            root / "questions.csv",
        ]
        found = [path for path in candidates if path.is_file()]
        if len(found) != 1:
            raise IntakeError("package must contain exactly one supported question workbook", field="questions")
        return found[0]

    @staticmethod
    def _validate_assessment_pools(lessons: list[dict], quiz: dict, questions: list[dict]) -> None:
        def candidates(selector: dict) -> int:
            wanted_objectives = {str(value) for value in selector.get("objective_codes") or []}
            wanted_tags = {str(value).casefold() for value in selector.get("tags_any") or []}
            count = 0
            for row in questions:
                row_tags = {tag.strip().casefold() for tag in str(row.get("tags") or "").split(",") if tag.strip()}
                if wanted_objectives and not wanted_objectives.intersection(
                    parse_objective_codes(row.get("objective_code"))
                ):
                    continue
                if wanted_tags and not row_tags.intersection(wanted_tags):
                    continue
                count += 1
            return count

        for lesson in lessons:
            quick = lesson.get("quick_check")
            if not quick:
                continue
            required = int(quick.get("displayed_count") or 3)
            if candidates(quick) < required:
                raise IntakeError(
                    f"Quick Check requests {required} questions but its filters match only {candidates(quick)}",
                    file="lessons",
                    field="quick_check",
                )
        blueprint = quiz.get("question_blueprint") or []
        requested = 0
        for entry in blueprint:
            required = int(entry.get("count") or 0)
            requested += required
            available = candidates(entry)
            if available < required:
                raise IntakeError(
                    f"Module Quiz blueprint requests {required} questions but its filters match only {available}",
                    file="module_quiz_blueprint.yaml",
                    field="question_blueprint",
                )
        displayed = int(quiz.get("displayed_count") or requested or 10)
        if displayed > len(questions) or (blueprint and requested != displayed):
            raise IntakeError(
                "Module Quiz displayed_count must equal its blueprint total and fit the bank",
                file="module_quiz_blueprint.yaml",
                field="displayed_count",
            )
        pass_percent = int(quiz.get("pass_percent") or 70)
        if not 1 <= pass_percent <= 100:
            raise IntakeError("pass_percent must be between 1 and 100", file="module_quiz_blueprint.yaml", field="pass_percent")
        categories = quiz.get("category_requirements") or []
        if categories:
            items = [
                {
                    "id": next(
                        (
                            tag
                            for tag in str(row.get("tags") or "").split(",")
                            if re.fullmatch(r"Q\d+", tag)
                        ),
                        str(index),
                    ),
                    "objective_codes": parse_objective_codes(row.get("objective_code")),
                    "tags": [tag for tag in str(row.get("tags") or "").split(",") if tag],
                }
                for index, row in enumerate(questions)
            ]
            try:
                selected = select_constrained(
                    items,
                    blueprint,
                    categories,
                    excluded_ids={
                        str(value) for value in quiz.get("exclude_question_ids") or []
                    },
                    rng=random.Random(0),
                )
            except ConstraintSelectionError as exc:
                raise IntakeError(
                    f"Module Quiz constraints are not satisfiable: {exc}",
                    file="module_quiz_blueprint.yaml",
                    field="category_requirements",
                ) from exc
            if len(selected) != displayed:
                raise IntakeError(
                    "Module Quiz objective quota total must equal displayed_count",
                    file="module_quiz_blueprint.yaml",
                    field="displayed_count",
                )

    def _normalize_resources(
        self,
        root: Path,
        version_key: str,
        lesson_keys: set[str],
        notes: set[str],
    ) -> list[dict]:
        path = root / "resources.yaml"
        if not path.is_file():
            return []
        rows = deepcopy(_read_yaml(path).get("resources") or [])
        for row in rows:
            if "type" in row and "resource_type" not in row:
                row["resource_type"] = row.pop("type")
            row["certification_version"] = version_key
            row["permission_status"] = _permission_status(
                row.get("permission_status"), notes
            )
            for link in row.get("links") or []:
                lesson_key = str(link.get("lesson_key") or "").strip()
                if lesson_key and lesson_key not in lesson_keys:
                    raise IntakeError("resource links to an unknown package lesson", file=path.name, field="lesson_key")
                if not lesson_key and not str(link.get("module_key") or "").strip():
                    raise IntakeError(
                        "resource link needs a lesson_key or module_key",
                        file=path.name,
                        field="links",
                    )
        return rows

    def _normalize_prompts(self, root: Path, version_key: str, module_key: str, domain: str, valid_objectives: set[str], notes: set[str]) -> list[dict]:
        path = root / "explain_prompts.yaml"
        if not path.is_file():
            return []
        doc = _read_yaml(path)
        rows = deepcopy(doc.get("prompts") or [])
        package_rubric_version = doc.get("rubric_version")
        for row in rows:
            if not row.get("prompt_key") and row.get("id"):
                row["prompt_key"] = (
                    f"interview.{_slug(module_key.removeprefix('module.')).replace('-', '.')}."
                    f"{_slug(str(row['id']))}"
                )
                notes.add(f"Explain {row['id']}: stable prompt_key generated from module key and id")
            if not row.get("rubric_version") and package_rubric_version:
                row["rubric_version"] = package_rubric_version
                notes.add(f"Explain {row['prompt_key']}: inherited package rubric_version")
            if not row.get("rubric") and row.get("minimum_for_pass") is not None:
                row["rubric"] = {
                    "minimum_for_pass": int(row["minimum_for_pass"]),
                    "partial_credit": doc.get("partial_credit"),
                }
                notes.add(f"Explain {row['prompt_key']}: minimum_for_pass -> rubric metadata")
            elif isinstance(row.get("rubric"), str) and row["rubric"].strip():
                row["rubric"] = {
                    "approved_text": row["rubric"].strip(),
                    "minimum_concepts": row.get("minimum_concepts"),
                    "must_include": deepcopy(row.get("must_include") or []),
                }
                notes.add(
                    f"Explain {row['prompt_key']}: plain-text rubric and explicit constraints -> rubric metadata"
                )
            row["certification_version"] = version_key
            row["module"] = module_key
            row["domain"] = domain
            row["importance"] = _normalize_scalar(row.get("importance", "working_knowledge"), IMPORTANCE_ALIASES, notes)
            self._validate_objectives([str(x) for x in row.get("objectives") or []], valid_objectives, path.name)
            for field in ("prompt_key", "prompt", "expected_concepts", "rubric", "rubric_version"):
                if row.get(field) in (None, "", []):
                    raise IntakeError(f"Explain field '{field}' is required", file=path.name, field=field)
        return rows

    @staticmethod
    def _normalize_practical(root: Path, notes: set[str]) -> tuple[list[dict], str | None]:
        path = root / "practical.yaml"
        if not path.is_file():
            return [], None
        doc = _read_yaml(path)
        rows = deepcopy(doc.get("labs") or ([doc["practical"]] if doc.get("practical") else []))
        if not rows and doc.get("title"):
            reserved = {
                "title",
                "objectives",
                "environment",
                "estimated_minutes",
                "purpose",
                "provenance",
            }
            exercise = {key: deepcopy(value) for key, value in doc.items() if key not in reserved}
            rows = [
                {
                    "title": doc["title"],
                    "description": doc.get("purpose"),
                    "estimated_minutes": doc.get("estimated_minutes"),
                    "environment_requirements": {"environment": doc.get("environment")},
                    "success_criteria": exercise,
                    "required_evidence": {
                        "student_tasks": [
                            task
                            for value in exercise.values()
                            if isinstance(value, dict)
                            for task in value.get("student_tasks") or []
                        ]
                    },
                    "source_name": doc.get("provenance"),
                }
            ]
            notes.add("top-level practical fields -> LabTemplate")
        if len(rows) != 1 or not str(rows[0].get("title") or "").strip():
            raise IntakeError("practical must define one titled LabTemplate", file=path.name, field="practical")
        rows[0].setdefault("lab_type", "guided")
        rows[0].setdefault("is_published", True)
        return rows, str(rows[0]["title"])

    def _service_desk(
        self,
        root: Path,
        module_key: str,
        valid_objectives: set[str],
        notes: set[str],
    ) -> tuple[str | None, list[dict]]:
        path = root / "service_desk.yaml"
        if not path.is_file():
            return None, []
        doc = _read_yaml(path)
        key = str(doc.get("scenario_key") or "").strip()
        if _has_inline_service_desk_content(doc):
            generated_key, scenario = _normalize_inline_service_desk(
                doc, module_key, valid_objectives, notes
            )
            return generated_key, [scenario]
        if self.service_desk_keys is not None and key not in self.service_desk_keys:
            raise IntakeError(f"Service Desk scenario '{key}' is not registered", file=path.name, field="scenario_key")
        return key, []

    @staticmethod
    def _build_assessments(module_key: str, lessons: list[dict], quiz: dict, quiz_title: str, practical_title: str | None, service_desk_key: str | None, has_prompts: bool) -> list[dict]:
        short = _slug(module_key.removeprefix("module."))
        rows = []
        for order, lesson in enumerate(sorted(lessons, key=lambda x: int(x.get("lesson_order") or 0)), 1):
            quick = lesson.get("quick_check")
            if not quick:
                continue
            config = {key: deepcopy(value) for key, value in quick.items() if key not in {"title", "displayed_count", "pass_percent"}}
            rows.append({
                "assessment_key": f"assess.{short}.qc.{_slug(str(lesson['lesson_key']).split('.')[-1])}",
                "assessment_role": "quick_check",
                "title": quick.get("title") or f"Quick Check — {lesson['title']}",
                "lesson_key": lesson["lesson_key"],
                "quiz_ref": quiz_title,
                "displayed_count": int(quick.get("displayed_count") or 3),
                "pass_percent": int(quick.get("pass_percent") or 60),
                "config": config,
                "display_order": order,
            })
        next_order = len(rows) + 1
        rows.append({
            "assessment_key": f"assess.{short}.module_quiz",
            "assessment_role": "module_quiz",
            "title": quiz.get("title") or f"Module Quiz — {quiz_title.removesuffix(' — Module Bank')}",
            "quiz_ref": quiz_title,
            "displayed_count": int(quiz.get("displayed_count") or 10),
            "pass_percent": int(quiz.get("pass_percent") or 70),
            "config": {
                "question_blueprint": deepcopy(quiz.get("question_blueprint") or []),
                "category_requirements": deepcopy(quiz.get("category_requirements") or []),
                "exclude_question_ids": deepcopy(quiz.get("exclude_question_ids") or []),
            },
            "display_order": next_order,
        })
        if practical_title:
            next_order += 1
            rows.append({"assessment_key": f"assess.{short}.practical", "assessment_role": "practical", "title": practical_title, "lab_ref": practical_title, "pass_percent": 70, "display_order": next_order})
        if service_desk_key:
            next_order += 1
            rows.append({"assessment_key": f"assess.{short}.service_desk", "assessment_role": "service_desk", "title": f"Service Desk — {service_desk_key}", "service_desk_ref": service_desk_key, "pass_percent": 70, "display_order": next_order})
        if has_prompts:
            next_order += 1
            rows.append({"assessment_key": f"assess.{short}.explain", "assessment_role": "explain", "title": "Explain", "pass_percent": 70, "display_order": next_order})
        return rows

    @staticmethod
    def _classify(module_root: Path, package_hash: str) -> str:
        if (module_root / package_hash / "IMPORT_RECEIPT.json").is_file():
            return "UNCHANGED"
        if module_root.is_dir() and any(path.is_dir() for path in module_root.iterdir()):
            return "CHANGED"
        return "NEW"

    def _write_runtime(self, normalized: dict, content: Path) -> list[str]:
        manifest = normalized["manifest"]
        cert_path = content / "certifications" / normalized["registry_path"]
        cert_doc = _read_yaml(cert_path)
        version = next(row for row in cert_doc["versions"] if row["version_key"] == manifest["version_key"])
        modules = version.setdefault("modules", [])
        existing = next((index for index, row in enumerate(modules) if row.get("module_key") == manifest["module_key"]), None)
        if existing is None:
            modules.append(normalized["module"])
        else:
            modules[existing] = normalized["module"]
        cert_path.write_text(yaml.safe_dump(cert_doc, sort_keys=False, allow_unicode=True), encoding="utf-8")

        slug = _slug(manifest["module_key"])
        lesson_root = content / "curriculum" / manifest["certification_key"] / manifest["version_key"] / slug
        # A changed package owns this module's runtime lesson directory.  Clear
        # the staged copy so removed lessons cannot survive as stale content.
        if lesson_root.exists():
            shutil.rmtree(lesson_root)
        lesson_root.mkdir(parents=True, exist_ok=True)
        for filename, rendered in normalized["lessons"]:
            (lesson_root / filename).write_text(rendered, encoding="utf-8")

        question_name = f"{slug}.csv"
        question_path = content / "questions" / question_name
        question_path.parent.mkdir(parents=True, exist_ok=True)
        with question_path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=TEMPLATE_COLUMNS)
            writer.writeheader()
            writer.writerows(normalized["questions"])
        approval_path = content / "questions" / APPROVALS_FILE
        approvals = _read_yaml(approval_path) if approval_path.exists() else {"approvals": []}
        approval = {
            "filename": question_name,
            "quiz_title": normalized["quiz_title"],
            "sha256": hashlib.sha256(question_path.read_bytes()).hexdigest(),
            "reviewed_question_count": manifest["question_count"],
            "editorial_status": "validated",
            "reviewed_at": datetime.now(timezone.utc).date().isoformat(),
            "review_note": f"External approved package {manifest['source_package_hash']}",
        }
        approvals["approvals"] = [row for row in approvals.get("approvals") or [] if row.get("filename") != question_name] + [approval]
        approval_path.write_text(yaml.safe_dump(approvals, sort_keys=False, allow_unicode=True), encoding="utf-8")

        destinations = [str(cert_path.relative_to(content)), str(lesson_root.relative_to(content)), str(question_path.relative_to(content)), str(approval_path.relative_to(content))]
        for folder, filename, key, rows in (
            ("resources", f"{slug}.yaml", "resources", normalized["resources"]),
            ("interview-prompts", f"{slug}.yaml", "prompts", normalized["prompts"]),
            ("labs", f"{slug}.yaml", "labs", normalized["labs"]),
            (
                "service-desk-scenarios",
                f"{slug}.yaml",
                "scenarios",
                normalized["service_desk_scenarios"],
            ),
        ):
            path = content / folder / filename
            if rows:
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_text(yaml.safe_dump({key: rows}, sort_keys=False, allow_unicode=True), encoding="utf-8")
            elif path.exists():
                path.unlink()
            # Include absent optional targets so apply removes stale files from
            # an earlier approved version of the same module.
            destinations.append(str(path.relative_to(content)))
        return destinations

    def _promote_runtime(self, staged: Path, relative_paths: list[str]) -> list[tuple[Path, bytes | None]]:
        rollback = []
        expanded: set[Path] = set()
        for relative in relative_paths:
            source = staged / relative
            if source.is_dir():
                expanded.update(path.relative_to(staged) for path in source.rglob("*") if path.is_file())
                destination_dir = self.content_dir / relative
                if destination_dir.is_dir():
                    expanded.update(path.relative_to(self.content_dir) for path in destination_dir.rglob("*") if path.is_file())
            else:
                expanded.add(Path(relative))
        try:
            for relative in sorted(expanded):
                source = staged / relative
                destination = self.content_dir / relative
                rollback.append((destination, destination.read_bytes() if destination.exists() else None))
                if not source.exists():
                    destination.unlink(missing_ok=True)
                    continue
                destination.parent.mkdir(parents=True, exist_ok=True)
                temporary = destination.with_name(f".{destination.name}.intake-{os.getpid()}")
                temporary.write_bytes(source.read_bytes())
                os.replace(temporary, destination)
        except Exception:
            self._rollback_runtime(rollback)
            raise
        return rollback

    @staticmethod
    def _rollback_runtime(rollback: list[tuple[Path, bytes | None]]) -> None:
        for destination, previous in reversed(rollback):
            if previous is None:
                destination.unlink(missing_ok=True)
            else:
                destination.write_bytes(previous)

    def _archive(self, source: Path, destination: Path, manifest: dict, loader_result: dict, runtime_destinations: list[str], classification: str) -> dict:
        destination.parent.mkdir(parents=True, exist_ok=True)
        temp = destination.parent / f".{destination.name}.pending-{os.getpid()}"
        if temp.exists():
            shutil.rmtree(temp)
        temp.mkdir()
        if source.is_file():
            shutil.copy2(source, temp / source.name)
        else:
            shutil.copytree(source, temp / "source")
        (temp / "NORMALIZED_MANIFEST.json").write_text(json.dumps(manifest, indent=2, sort_keys=True) + "\n", encoding="utf-8")
        count_source = loader_result.get("first", loader_result)
        totals = {
            key: int(count_source.get(key, 0)) if isinstance(count_source, dict) else 0
            for key in ("created", "updated", "unchanged")
        }
        receipt = {
            "timestamp": datetime.now(timezone.utc).isoformat(),
            "package_hash": manifest["source_package_hash"],
            "original_filename": manifest["original_filename"],
            "certification": manifest["certification_key"],
            "version": manifest["version_key"],
            "domain": manifest["domain_key"],
            "module_key": manifest["module_key"],
            "module_title": manifest["module_title"],
            "lesson_count": manifest["lesson_count"],
            "question_count": manifest["question_count"],
            "resource_count": manifest["resource_count"],
            "practical_present": manifest["practical_present"],
            "service_desk_present": manifest["service_desk_present"],
            "explain_count": manifest["explain_count"],
            "objective_codes": manifest["objectives"],
            "approved_destination": str(destination),
            "runtime_destinations": runtime_destinations,
            "loader_importer_results": loader_result,
            **totals,
            "classification": classification,
            "validation_result": "VALID",
            "tool_version": "nexus-curriculum-intake/1",
        }
        (temp / "IMPORT_RECEIPT.json").write_text(json.dumps(receipt, indent=2, sort_keys=True) + "\n", encoding="utf-8")
        os.replace(temp, destination)
        return receipt

    @staticmethod
    def _verify_archive(destination: Path, package_hash: str) -> None:
        receipt_path = destination / "IMPORT_RECEIPT.json"
        if not receipt_path.is_file():
            raise IntakeError("approved archive receipt is missing", field="receipt")
        receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
        if receipt.get("package_hash") != package_hash:
            raise IntakeError("approved archive hash does not match inbox package", field="package_hash")
        source_dir = destination / "source"
        if source_dir.is_dir():
            archived_hash = _tree_hash(source_dir)
        else:
            archives = [path for path in destination.glob("*.zip") if path.is_file()]
            if len(archives) != 1:
                raise IntakeError("approved archive does not contain exactly one preserved source", field="approved_source")
            archived_hash = hashlib.sha256(archives[0].read_bytes()).hexdigest()
        if archived_hash != package_hash:
            raise IntakeError("preserved approved source failed its SHA-256 verification", field="package_hash")

    def _remove_inbox_item(self, source: Path) -> None:
        if not _within(source, self.dropbox_dir) or source.parent.resolve() != self.dropbox_dir.resolve():
            raise IntakeError("refusing to remove an item outside the dropbox", field="cleanup")
        if source.is_dir():
            shutil.rmtree(source)
        else:
            source.unlink()

    def _write_report(self, source_name: str, result: dict) -> None:
        report_dir = self.dropbox_dir / ".reports"
        report_dir.mkdir(parents=True, exist_ok=True)
        (report_dir / f"{_slug(source_name) or 'package'}.json").write_text(json.dumps(result, indent=2, sort_keys=True) + "\n", encoding="utf-8")

    def _validate_in_scratch_database(self, content: Path, manifest: dict) -> dict:
        """Run real V2 loaders twice against a fresh migrated disposable DB."""
        with tempfile.TemporaryDirectory(prefix="nexus-curriculum-db-") as tmp:
            database = Path(tmp) / "intake.db"
            env = os.environ.copy()
            env["DATABASE_URL"] = f"sqlite:///{database}"
            env["NEXUS_CONTENT_DIR"] = str(content)
            env["NEXUS_MODULE_KEY"] = manifest["module_key"]
            subprocess.run(
                [sys.executable, "-m", "alembic", "upgrade", "head"],
                cwd=self.backend_dir,
                env=env,
                check=True,
                capture_output=True,
                text=True,
            )
            script = """
import json, os
from app.database import SessionLocal
from app.models.certification import CertificationModule, ModuleAssessment
from app.services.v2_content_loader import load_module
from seed import seed_service_desk_scenarios
root = os.environ['NEXUS_CONTENT_DIR']
kw = {name: os.path.join(root, folder) for name, folder in {
 'cert_dir':'certifications','objectives_dir':'objectives','curriculum_dir':'curriculum',
 'resources_dir':'resources','interview_prompts_dir':'interview-prompts',
 'questions_dir':'questions','labs_dir':'labs',
 'service_desk_scenarios_dir':'service-desk-scenarios'}.items()}
db = SessionLocal()
try:
 seed_service_desk_scenarios(db)
 db.commit()
 first = load_module(db, commit=True, **kw)
 second = load_module(db, commit=True, **kw)
 if second['created'] or second['updated']:
  raise RuntimeError(f"second V2 loader pass was not idempotent: {second}")
 module = db.query(CertificationModule).filter_by(module_key=os.environ['NEXUS_MODULE_KEY']).one()
 package_assessments = db.query(ModuleAssessment).filter_by(certification_module_id=module.id).all()
 content_roles = {'quick_check', 'module_quiz', 'practical'}
 sd = db.query(ModuleAssessment).filter_by(certification_module_id=module.id, assessment_role='service_desk').all()
 references = dict(second['references'])
 references['package_content_engine_unresolved'] = [
  row.assessment_key for row in package_assessments
  if row.assessment_role in content_roles and not (row.quiz_id or row.lab_template_id)
 ]
 references['package_service_desk_unresolved'] = [row.assessment_key for row in sd if not row.service_desk_scenario_id]
 print(json.dumps({'first': first, 'second': second, 'references': references}))
finally:
 db.close()
"""
            try:
                completed = subprocess.run(
                    [sys.executable, "-c", script], cwd=self.backend_dir, env=env,
                    check=True, capture_output=True, text=True,
                )
            except subprocess.CalledProcessError as exc:
                detail = (exc.stderr or exc.stdout or str(exc)).strip()
                raise IntakeError(
                    f"scratch V2 loading failed: {detail}",
                    field="runtime_validation",
                ) from exc
            return json.loads(completed.stdout.strip().splitlines()[-1])
